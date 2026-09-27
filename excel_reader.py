"""
excel_reader.py
----------------
Reads recipient rows from an .xlsx workbook and turns them into validated
Recipient objects. Rows with missing/malformed emails are skipped (and
reported) rather than silently sent to or silently crashing the whole run.

Expected columns (case-insensitive, order-independent):
    SNo, Name, Email, Title, Company

Only "Name" and "Email" are strictly required; "Title" and "Company" are
used for personalization if present, and default to "" otherwise.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import List

import openpyxl

EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")


class ExcelReadError(RuntimeError):
    """Raised when the workbook can't be read or has no usable header row."""


@dataclass(frozen=True)
class Recipient:
    name: str
    email: str
    title: str = ""
    company: str = ""
    row_number: int = -1  # 1-based row in the source sheet, for error messages

    @property
    def first_name(self) -> str:
        return self.name.strip().split(" ")[0] if self.name.strip() else self.name


def _normalize_header(cell_value) -> str:
    return str(cell_value or "").strip().lower()


def load_recipients(excel_path: Path) -> tuple[List[Recipient], List[str]]:
    """
    Returns (recipients, warnings).

    `warnings` contains human-readable strings describing any row that was
    skipped (missing name/email, invalid email format, etc.) so the caller
    can surface them instead of silently dropping people.
    """
    if not Path(excel_path).exists():
        raise ExcelReadError(f"Excel file not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    header_row = None
    header_row_idx = None
    for idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=10), start=1):
        values = [_normalize_header(c.value) for c in row]
        if "email" in values and "name" in values:
            header_row = values
            header_row_idx = idx
            break

    if header_row is None:
        raise ExcelReadError(
            "Could not find a header row containing both 'Name' and 'Email' "
            "columns in the first 10 rows of the sheet."
        )

    col_index = {name: i for i, name in enumerate(header_row)}

    recipients: List[Recipient] = []
    warnings: List[str] = []

    for row in sheet.iter_rows(min_row=header_row_idx + 1):
        row_num = row[0].row
        if all(c.value in (None, "") for c in row):
            continue  # blank row

        def get(col_name: str) -> str:
            idx = col_index.get(col_name)
            if idx is None or idx >= len(row):
                return ""
            val = row[idx].value
            return str(val).strip() if val is not None else ""

        name = get("name")
        email = get("email")
        title = get("title")
        company = get("company")

        if not name or not email:
            warnings.append(f"Row {row_num}: skipped — missing name or email.")
            continue

        if not EMAIL_RE.match(email):
            warnings.append(f"Row {row_num}: skipped — invalid email format: {email!r}")
            continue

        recipients.append(
            Recipient(
                name=name, email=email.lower(), title=title, company=company,
                row_number=row_num,
            )
        )

    # De-duplicate by email within this same file, keeping the first occurrence.
    seen = set()
    deduped: List[Recipient] = []
    for r in recipients:
        if r.email in seen:
            warnings.append(f"Row {r.row_number}: skipped — duplicate email in sheet: {r.email}")
            continue
        seen.add(r.email)
        deduped.append(r)

    return deduped, warnings
